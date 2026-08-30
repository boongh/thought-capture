from __future__ import annotations

import hashlib
import re
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "DESIGN.md"
OUTPUT = ROOT / "docs" / "Thought-Capture-AI-System-Design.docx"
QA_DIR = ROOT / "rendered"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17324D"
MUTED = "667085"
LIGHT = "F2F4F7"
CALLOUT = "EAF2F8"
WHITE = "FFFFFF"
BLACK = "111827"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D0D5DD", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def apply_table_geometry(table, widths: list[int], indent=120) -> None:
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Update this field in Word to refresh the table of contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def define_numbering(document: Document) -> tuple[int, int]:
    numbering = document.part.numbering_part.element

    def make_num(abstract_id: int, num_id: int, fmt: str, text: str, font: str | None = None) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.extend([tabs, ind])
        lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
        if font:
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font)
            fonts.set(qn("w:hAnsi"), font)
            r_pr.append(fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)

    existing = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    base = max(existing, default=20) + 10
    make_num(base, base, "bullet", "•", "Calibri")
    make_num(base + 1, base + 1, "decimal", "%1.")
    return base, base + 1


def set_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def add_inline_runs(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|https?://\S+)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
            if part.startswith("http"):
                run.font.color.rgb = RGBColor.from_string(BLUE)
                run.underline = True


def add_code_block(doc: Document, language: str, lines: list[str]) -> None:
    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(4)
    label.paragraph_format.space_after = Pt(2)
    label_run = label.add_run(language.upper() if language else "CODE")
    label_run.bold = True
    label_run.font.size = Pt(8)
    label_run.font.color.rgb = RGBColor.from_string(MUTED)
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.08)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F7F8FA")
        p_pr.append(shd)
        run = p.add_run(line or " ")
        run.font.name = "Consolas"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(BLACK)
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(4)


def diagram_font(size: int, bold: bool = False):
    name = "arialbd.ttf" if bold else "arial.ttf"
    path = Path("C:/Windows/Fonts") / name
    return ImageFont.truetype(str(path), size) if path.exists() else ImageFont.load_default()


def draw_centered(draw, box, text, font, fill=BLACK, max_chars=24):
    x1, y1, x2, y2 = box
    wrapped = "\n".join(textwrap.wrap(text, max_chars))
    bounds = draw.multiline_textbbox((0, 0), wrapped, font=font, spacing=5, align="center")
    w, h = bounds[2] - bounds[0], bounds[3] - bounds[1]
    color = fill if str(fill).startswith("#") else f"#{fill}"
    draw.multiline_text(((x1 + x2 - w) / 2, (y1 + y2 - h) / 2), wrapped, font=font, fill=color, spacing=5, align="center")


def rounded_box(draw, box, title, subtitle="", fill="EAF2F8"):
    draw.rounded_rectangle(box, radius=18, fill=f"#{fill}", outline=f"#{BLUE}", width=3)
    x1, y1, x2, y2 = box
    draw.text((x1 + 18, y1 + 13), title, font=diagram_font(23, True), fill=f"#{NAVY}")
    if subtitle:
        wrapped = "\n".join(textwrap.wrap(subtitle, 27))
        draw.multiline_text((x1 + 18, y1 + 50), wrapped, font=diagram_font(17), fill=f"#{BLACK}", spacing=4)


def arrow(draw, start, end, color=NAVY, width=4):
    draw.line([start, end], fill=f"#{color}", width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        sign = 1 if ex > sx else -1
        points = [(ex, ey), (ex - sign * 15, ey - 9), (ex - sign * 15, ey + 9)]
    else:
        sign = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 9, ey - sign * 15), (ex + 9, ey - sign * 15)]
    draw.polygon(points, fill=f"#{color}")


def create_architecture_diagram(path: Path) -> None:
    img = Image.new("RGB", (1500, 850), "white")
    d = ImageDraw.Draw(img)
    d.text((40, 24), "Runtime architecture", font=diagram_font(34, True), fill=f"#{NAVY}")
    d.text((40, 68), "One canonical store, two complementary retrieval paths", font=diagram_font(21), fill=f"#{MUTED}")
    rounded_box(d, (40, 155, 260, 275), "Discord", "capture and digest", "F2F4F7")
    rounded_box(d, (330, 155, 570, 275), "Bot", "allowlist, ingest adapter")
    rounded_box(d, (650, 155, 930, 275), "Application", "use cases and API gateway")
    rounded_box(d, (1010, 110, 1450, 250), "PostgreSQL", "raw log, revisions, entities, exact search", "E8F5E9")
    rounded_box(d, (1010, 330, 1450, 470), "Khoj", "semantic index, reranking, Ask/RAG", "FFF4E5")
    rounded_box(d, (650, 520, 930, 660), "Worker", "20:00 organize, export, retry")
    rounded_box(d, (1010, 570, 1450, 710), "OpenRouter", "structured completions; provider gateway", "FCEEF5")
    rounded_box(d, (40, 520, 570, 660), "Interface", "Khoj UI now; unified custom UI later", "F2F4F7")
    arrow(d, (260, 215), (330, 215))
    arrow(d, (570, 215), (650, 215))
    arrow(d, (930, 190), (1010, 180))
    arrow(d, (930, 230), (1010, 390))
    arrow(d, (790, 520), (790, 275))
    arrow(d, (930, 590), (1010, 640))
    arrow(d, (930, 555), (1100, 470))
    arrow(d, (570, 590), (650, 590))
    arrow(d, (300, 520), (740, 275))
    d.text((1080, 275), "deterministic filters", font=diagram_font(17, True), fill=f"#{MUTED}")
    d.text((1080, 500), "semantic path", font=diagram_font(17, True), fill=f"#{MUTED}")
    img.save(path, dpi=(180, 180))


def create_capture_sequence(path: Path) -> None:
    img = Image.new("RGB", (1500, 850), "white")
    d = ImageDraw.Draw(img)
    d.text((40, 24), "Durable capture sequence", font=diagram_font(34, True), fill=f"#{NAVY}")
    labels = ["Discord user", "Bot", "Capture use case", "PostgreSQL", "Attachment store"]
    xs = [120, 400, 730, 1060, 1370]
    for x, label in zip(xs, labels):
        d.rounded_rectangle((x - 105, 110, x + 105, 180), radius=14, fill="#EAF2F8", outline=f"#{BLUE}", width=3)
        draw_centered(d, (x - 100, 112, x + 100, 178), label, diagram_font(20, True), max_chars=18)
        d.line((x, 180, x, 790), fill="#98A2B3", width=2)
    events = [
        (120, 400, 250, "message + files"),
        (400, 730, 340, "normalized command"),
        (730, 1370, 430, "validate, hash, atomic store"),
        (730, 1060, 525, "transaction: thought + links + outbox"),
        (1060, 730, 615, "thought ID / existing ID"),
        (730, 400, 695, "durable result"),
        (400, 120, 765, "acknowledgement"),
    ]
    for x1, x2, y, label in events:
        arrow(d, (x1, y), (x2, y), color=BLUE, width=4)
        bounds = d.textbbox((0, 0), label, font=diagram_font(17))
        tw = bounds[2] - bounds[0]
        d.rectangle(((x1 + x2 - tw) / 2 - 8, y - 29, (x1 + x2 + tw) / 2 + 8, y - 6), fill="white")
        d.text(((x1 + x2 - tw) / 2, y - 29), label, font=diagram_font(17), fill=f"#{BLACK}")
    img.save(path, dpi=(180, 180))


def set_picture_alt(inline_shape, alt_text: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", alt_text)


def add_mermaid_diagram(doc: Document, lines: list[str]) -> None:
    QA_DIR.mkdir(parents=True, exist_ok=True)
    code = "\n".join(lines)
    if "sequenceDiagram" in code:
        path = QA_DIR / "capture-sequence.png"
        create_capture_sequence(path)
        alt = "Sequence diagram showing Discord capture committed to attachment storage and PostgreSQL before acknowledgement."
        caption = "Figure 2. Durable capture: acknowledgement follows the canonical database transaction."
    else:
        path = QA_DIR / "runtime-architecture.png"
        create_architecture_diagram(path)
        alt = "Architecture diagram showing Discord, first-party services, PostgreSQL, Khoj, OpenRouter, and current and future interfaces."
        caption = "Figure 1. PostgreSQL is canonical and precise; Khoj supplies semantic retrieval and Ask."
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(path), width=Inches(6.35))
    set_picture_alt(shape, alt)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(8)
    cr = cp.add_run(caption)
    cr.italic = True
    cr.font.size = Pt(9)
    cr.font.color.rgb = RGBColor.from_string(MUTED)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        values = [v.strip() for v in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", v) for v in values):
            rows.append(values)
        i += 1
    return rows, i


def table_widths(rows: list[list[str]]) -> list[int]:
    cols = max(len(r) for r in rows)
    max_len = []
    for c in range(cols):
        max_len.append(max(6, min(80, max((len(r[c]) if c < len(r) else 0) for r in rows))))
    total_weight = sum(max_len)
    raw = [max(900, round(9360 * n / total_weight)) for n in max_len]
    scale = 9360 / sum(raw)
    widths = [max(800, round(v * scale)) for v in raw]
    widths[-1] += 9360 - sum(widths)
    minimum = 1100 if cols <= 3 else 850
    for idx in range(cols):
        if widths[idx] < minimum:
            needed = minimum - widths[idx]
            donor = max((j for j in range(cols) if j != idx), key=lambda j: widths[j])
            widths[idx] += needed
            widths[donor] -= needed
    return widths


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    widths = table_widths(rows)
    apply_table_geometry(table, widths)
    set_table_borders(table)
    for r_idx, values in enumerate(rows):
        row = table.rows[r_idx]
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = values[c_idx] if c_idx < len(values) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            add_inline_runs(p, text)
            for run in p.runs:
                run.font.size = Pt(8.5 if cols >= 4 else 9)
                if r_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(NAVY)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT)
        if r_idx == 0:
            set_repeat_table_header(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
        "Heading 4": (11, NAVY, 7, 3),
    }
    for name, (size, color, before, after) in tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def add_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("THOUGHT CAPTURE AI  /  SYSTEM DESIGN 1.0")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Internal design anchor  |  ")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_field(fp, "PAGE")
    fp.add_run(" of ")
    add_page_field(fp, "NUMPAGES")


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("SYSTEM DESIGN")
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    tr = title.add_run("Thought Capture AI")
    tr.font.name = "Calibri"
    tr.font.size = Pt(30)
    tr.font.bold = True
    tr.font.color.rgb = RGBColor.from_string(NAVY)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(26)
    sr = subtitle.add_run("A durable, reversible personal memory system built around Discord, PostgreSQL, Khoj, and OpenRouter")
    sr.font.size = Pt(14)
    sr.font.color.rgb = RGBColor.from_string(MUTED)

    values = [
        ("Status", "Accepted implementation anchor"),
        ("Version", "1.0"),
        ("Date", "30 August 2026"),
        ("Audience", "Small experienced engineering team"),
        ("Primary outcome", "Discord capture to cited 20:00 digest and hybrid retrieval"),
    ]
    for label, value in values:
        meta = doc.add_paragraph()
        meta.paragraph_format.space_before = Pt(0)
        meta.paragraph_format.space_after = Pt(3)
        lr = meta.add_run(f"{label}: ")
        lr.bold = True
        lr.font.size = Pt(10.5)
        lr.font.color.rgb = RGBColor.from_string(NAVY)
        vr = meta.add_run(value)
        vr.font.size = Pt(10.5)
        vr.font.color.rgb = RGBColor.from_string(BLACK)

    callout = doc.add_paragraph()
    callout.paragraph_format.space_before = Pt(28)
    callout.paragraph_format.space_after = Pt(6)
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.18)
    p_pr = callout._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CALLOUT)
    p_pr.append(shd)
    cr = callout.add_run("Decision: PostgreSQL remains canonical and precise; Khoj provides semantic retrieval and Ask. A stable gateway makes a future unified UI straightforward without maintaining a Khoj fork.")
    cr.bold = True
    cr.font.color.rgb = RGBColor.from_string(NAVY)
    cr.font.size = Pt(10.5)

    doc.add_page_break()
    h = doc.add_paragraph("Contents", style="Heading 1")
    h.paragraph_format.space_before = Pt(0)
    toc = doc.add_paragraph()
    add_toc(toc)
    note = doc.add_paragraph("The document also uses Word heading styles for navigation and accessibility.")
    note.runs[0].italic = True
    note.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_page_break()


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    add_header_footer(doc.sections[0])
    bullet_id, decimal_id = define_numbering(doc)
    add_cover(doc)

    i = 0
    in_code = False
    code_language = ""
    code_lines: list[str] = []
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            p = doc.add_paragraph()
            add_inline_runs(p, " ".join(x.strip() for x in paragraph_buffer))
            paragraph_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_paragraph()
            if not in_code:
                in_code = True
                code_language = stripped[3:].strip()
                code_lines = []
            else:
                if code_language.lower() == "mermaid":
                    add_mermaid_diagram(doc, code_lines)
                else:
                    add_code_block(doc, code_language, code_lines)
                in_code = False
                code_language = ""
                code_lines = []
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not stripped:
            flush_paragraph()
            i += 1
            continue
        if stripped.startswith("# "):
            # The cover already carries the document title.
            flush_paragraph()
            i += 1
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = min(len(heading.group(1)) - 1, 3)
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline_runs(p, heading.group(2))
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            flush_paragraph()
            rows, i = parse_table(lines, i)
            if rows:
                add_markdown_table(doc, rows)
            continue
        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            flush_paragraph()
            p = doc.add_paragraph()
            set_num(p, bullet_id)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.167
            add_inline_runs(p, bullet.group(1))
            i += 1
            continue
        number = re.match(r"^\d+\.\s+(.+)$", stripped)
        if number:
            flush_paragraph()
            p = doc.add_paragraph()
            set_num(p, decimal_id)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.167
            add_inline_runs(p, number.group(1))
            i += 1
            continue
        checkbox = re.match(r"^- \[([ xX])\]\s+(.+)$", stripped)
        if checkbox:
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            add_inline_runs(p, ("☒ " if checkbox.group(1).lower() == "x" else "☐ ") + checkbox.group(2))
            i += 1
            continue
        if stripped.startswith("**") and stripped.endswith("**") and len(stripped) < 180:
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            add_inline_runs(p, stripped)
            i += 1
            continue
        paragraph_buffer.append(line)
        i += 1
    flush_paragraph()

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    core = doc.core_properties
    core.title = "Thought Capture AI - System Design"
    core.subject = "Accepted implementation anchor"
    core.author = "Thought Capture AI project"
    core.keywords = "Discord, Khoj, OpenRouter, PostgreSQL, personal memory, system design"
    core.comments = f"Generated from docs/DESIGN.md sha256={hashlib.sha256(SOURCE.read_bytes()).hexdigest()}"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
