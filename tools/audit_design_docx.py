from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "Thought-Capture-AI-System-Design.docx"


def main() -> None:
    doc = Document(DOCX)
    text = "\n".join(p.text for p in doc.paragraphs)
    required = [
        "Executive decision",
        "Release 1 definition of done",
        "Khoj integration contract",
        "OpenRouter for a new AI developer",
        "Privacy, security, and threat model",
        "Delivery plan and gates",
        "Release 1 acceptance checklist",
    ]
    missing = [item for item in required if item not in text]
    assert not missing, f"missing required sections: {missing}"
    assert len(doc.paragraphs) > 500
    assert len(doc.tables) >= 5
    headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert len(headings) >= 60
    section = doc.sections[0]
    assert round(section.page_width.inches, 2) == 8.50
    assert round(section.page_height.inches, 2) == 11.00
    assert all(
        round(x.inches, 2) == 1.00
        for x in (
            section.left_margin,
            section.right_margin,
            section.top_margin,
            section.bottom_margin,
        )
    )

    numbering = doc.part.numbering_part.element
    assert len(numbering.findall(qn("w:num"))) >= 2
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        assert tbl_pr.find(qn("w:tblW")).get(qn("w:w")) == "9360"
        assert tbl_pr.find(qn("w:tblInd")).get(qn("w:w")) == "120"

    drawings = doc.part.element.findall(".//" + qn("w:drawing"))
    assert len(drawings) == 2, f"expected 2 diagrams, found {len(drawings)}"
    doc_prs = doc.part.element.findall(".//" + qn("wp:docPr"))
    assert all(node.get("descr") for node in doc_prs), "every diagram needs alt text"
    print(
        f"PASS paragraphs={len(doc.paragraphs)} headings={len(headings)} tables={len(doc.tables)} diagrams={len(drawings)}"
    )


if __name__ == "__main__":
    main()
